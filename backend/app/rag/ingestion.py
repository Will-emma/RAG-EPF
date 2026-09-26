from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation
from langchain_text_splitters import RecursiveCharacterTextSplitter
from fastembed import TextEmbedding

from app.core.config import settings

# fastembed exécute le même modèle all-MiniLM-L6-v2 (384 dimensions, vecteurs
# identiques) via ONNX Runtime, sans torch : ~250 Mo de RAM au lieu de ~900 Mo,
# ce qui permet d'héberger le backend sur une petite instance.
_embedding_model: TextEmbedding | None = None

# Nombre de passages encodés à la fois. Par défaut fastembed en encode 256 :
# pour des passages de 800 caractères, la mémoire de travail dépasse alors les
# 512 Mo de l'instance Render (arrêt forcé en plein import d'un gros cours).
# Par paquets de 8 : ~250 Mo au pic, même pour 1 500 passages, sans perte de vitesse.
EMBEDDING_BATCH_SIZE = 8


def get_embedding_model() -> TextEmbedding:
    global _embedding_model
    if _embedding_model is None:
        _embedding_model = TextEmbedding(
            settings.EMBEDDING_MODEL, cache_dir=settings.EMBEDDING_CACHE_DIR
        )
    return _embedding_model


def extract_pages(file_path: Path) -> list[tuple[int, str]]:
    """Renvoie une liste de (numéro_de_page, texte) selon le type de fichier."""
    ext = file_path.suffix.lower()

    if ext == ".pdf":
        pages = []
        with fitz.open(file_path) as doc:
            for i, page in enumerate(doc, start=1):
                text = page.get_text()
                if text.strip():
                    pages.append((i, text))
        return pages

    if ext == ".pptx":
        prs = Presentation(file_path)
        pages = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame
            ]
            content = "\n".join(t for t in texts if t.strip())
            if content.strip():
                pages.append((i, content))
        return pages

    if ext == ".docx":
        doc = DocxDocument(file_path)
        content = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [(1, content)] if content.strip() else []

    raise ValueError(f"Format non supporté: {ext}")


def chunk_pages(
    pages: list[tuple[int, str]], chunk_size: int = 800, chunk_overlap: int = 100
) -> list[tuple[int, str]]:
    """Découpe chaque page en chunks, en conservant le numéro de page."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size, chunk_overlap=chunk_overlap
    )
    chunks = []
    for page_number, text in pages:
        for piece in splitter.split_text(text):
            chunks.append((page_number, piece))
    return chunks


def embed_texts(texts: list[str]) -> list[list[float]]:
    model = get_embedding_model()
    return [
        embedding.tolist()
        for embedding in model.embed(texts, batch_size=EMBEDDING_BATCH_SIZE)
    ]
